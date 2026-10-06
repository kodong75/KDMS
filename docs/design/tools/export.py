"""배포본 만들기: build.py 결과에서 PNG(미리보기), PDF(문서), PPTX(발표용 한 장).

    python3 docs/design/tools/build.py
    python3 docs/design/tools/export.py

필요한 것
- Chromium 계열 브라우저(헤드리스). 환경 변수 CHROME 로 경로를 지정할 수 있다.
  창 크기를 정확히 지키는 headless_shell(Playwright 배포본)을 먼저 찾는다.
- PPTX 는 python-pptx(MIT). 없으면 PPTX 만 건너뛴다.
"""

from __future__ import annotations

import glob
import os
import shutil
import subprocess
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
import build  # noqa: E402
from diagram import DESIGN_DIR  # noqa: E402

DIST = DESIGN_DIR / "dist"
SCALE = 2  # PNG 배율(1600x900 → 3200x1800)


def find_chrome() -> str:
    if os.environ.get("CHROME"):
        return os.environ["CHROME"]
    candidates = sorted(glob.glob("/opt/pw-browsers/chromium_headless_shell-*/*/headless_shell"))
    candidates += sorted(glob.glob(os.path.expanduser("~/Library/Caches/ms-playwright/chromium_headless_shell-*/*/headless_shell")))
    candidates += ["/Applications/Google Chrome.app/Contents/MacOS/Google Chrome"]
    candidates += [shutil.which(n) or "" for n in ("chromium", "chromium-browser", "google-chrome")]
    for c in candidates:
        if c and Path(c).exists():
            return c
    sys.exit("Chromium 을 찾지 못했다. CHROME=<경로> 로 지정한다.")


def chrome(args: list[str]) -> None:
    exe = find_chrome()
    flags = ["--no-sandbox", "--hide-scrollbars", "--virtual-time-budget=4000"]
    if not exe.endswith("headless_shell"):
        flags.insert(0, "--headless")
    subprocess.run([exe, *flags, *args], check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)


def export_doc(doc_id: str) -> None:
    svg = DESIGN_DIR / "diagrams" / f"{doc_id}.svg"
    html = DESIGN_DIR / f"{doc_id}.html"
    png = DIST / f"{doc_id}.png"
    pdf = DIST / f"{doc_id}.pdf"
    chrome([f"--screenshot={png}", "--window-size=1600,900", f"--force-device-scale-factor={SCALE}", svg.as_uri()])
    chrome(["--no-pdf-header-footer", f"--print-to-pdf={pdf}", html.as_uri()])
    try:
        from pptx import Presentation
        from pptx.util import Emu
    except ImportError:
        print("python-pptx 없음: PPTX 건너뜀")
        return
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(12192000), Emu(6858000)  # 16:9 (33.867 x 19.05 cm)
    slide = prs.slides.add_slide(prs.slide_layouts[6])               # 빈 레이아웃
    slide.shapes.add_picture(str(png), 0, 0, prs.slide_width, prs.slide_height)
    prs.core_properties.title = f"KDMS {doc_id}"
    prs.core_properties.author = "KDMS"
    prs.save(DIST / f"{doc_id}.pptx")


def main() -> None:
    build.main()
    DIST.mkdir(exist_ok=True)
    for doc in build.DOCS:
        if "id" in doc:
            export_doc(doc["id"])
            print("exported:", doc["id"], "(png, pdf, pptx)")


if __name__ == "__main__":
    main()
