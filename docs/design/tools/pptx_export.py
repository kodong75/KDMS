"""다이어그램 모델 → 편집 가능한 PowerPoint(python-pptx, MIT). 그림 한 장 = 슬라이드 하나.

그림을 이미지로 넣지 않는다. 영역·블록은 둥근 사각형(글은 도형 안), 아이콘은 자유형 도형,
선은 블록에 붙인 연결선(꺾인·직선 연결선)이라 블록을 옮기면 PowerPoint 가 선을 다시 잇는다.
ERD 엔터티·관계선은 도형·글상자를 묶은 그룹이다(엔터티 하나 = 그룹 하나, 컬럼 글은 열마다 글상자).
좌표는 diagram.py 의 1600x900 캔버스를 16:9 슬라이드에 그대로 옮긴다.
"""

from __future__ import annotations

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

from diagram import C, E_HEAD, E_ROW, FONT_PPTX, LINE_KINDS, T, Diagram, Entity, Rel, ie_marks, text_width, icon_paths
from svgpath import flatten

SLIDE_W, SLIDE_H = 12192000, 6858000      # 16:9
PX = SLIDE_W / 1600                        # EMU / 캔버스 px
PT = PX / 12700                            # pt / 캔버스 px
CXN = {"t": 0, "l": 1, "b": 2, "r": 3}     # 사각형·마름모 연결점 번호
CXN_OVAL = {"t": 0, "l": 2, "b": 4, "r": 6}  # 타원(연결점 원)은 8개


def E(v: float) -> Emu:
    return Emu(int(round(v * PX)))


def rgb(hex_: str) -> RGBColor:
    return RGBColor.from_string(hex_.lstrip("#"))


def _plain(shape) -> None:
    """테마 스타일(그림자 등)을 떼어 플랫하게."""
    style = shape._element.find(qn("p:style"))
    if style is not None:
        shape._element.remove(style)


def _font(run, size_px: float, bold: bool, color: str) -> None:
    f = run.font
    f.size = Pt(size_px * PT)
    f.bold = bold
    f.color.rgb = rgb(color)
    f.name = FONT_PPTX
    rpr = run._r.get_or_add_rPr()
    for tag in ("a:ea", "a:cs"):
        el = rpr.find(qn(tag))
        if el is None:
            el = etree.SubElement(rpr, qn(tag))
        el.set("typeface", FONT_PPTX)


def _paras(tf, lines: list[tuple[str, float, bool, str, float | None]], align, wrap=False) -> None:
    """lines: (글, 크기px, 굵게, 색, 줄높이px)."""
    tf.word_wrap = wrap
    tf.auto_size = MSO_AUTO_SIZE.NONE
    for i, (text, size, bold, color, lh) in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        if lh:
            p.line_spacing = Pt(lh * PT)
        r = p.add_run()
        r.text = text
        _font(r, size, bold, color)


def _margins(tf, l=0.0, t=0.0, r=0.0, b=0.0) -> None:
    tf.margin_left, tf.margin_top, tf.margin_right, tf.margin_bottom = E(l), E(t), E(r), E(b)


def rrect(slide, x, y, w, h, fill, line, lw, radius, name):
    s = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, E(x), E(y), E(w), E(h))
    _plain(s)
    s.adjustments[0] = min(0.5, radius / min(w, h))
    s.fill.solid()
    s.fill.fore_color.rgb = rgb(fill)
    if line:
        s.line.color.rgb = rgb(line)
        s.line.width = Pt(lw * PT)
    else:
        s.line.fill.background()
    s.name = name
    return s


def textbox(slide, x, y, w, h, lines, align=PP_ALIGN.LEFT, name="글", fill=None):
    tb = slide.shapes.add_textbox(E(x), E(y), E(w), E(h))
    tf = tb.text_frame
    _margins(tf)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    _paras(tf, lines, align, wrap=True)   # 폭 고정(자동 폭이면 뷰어마다 가운데로 몰린다)
    if fill:
        tb.fill.solid()
        tb.fill.fore_color.rgb = rgb(fill)
    tb.name = name
    return tb


def icon(slide, name, x, y, size, color):
    """Material 아이콘을 자유형 도형으로(viewBox 0 -960 960 960)."""
    polys = [p for d in icon_paths(name) for p in flatten(d)]
    k = size / 960.0
    first = polys[0][0]
    fb = slide.shapes.build_freeform(E(x + first[0] * k), E(y + (first[1] + 960) * k), scale=1.0)
    for j, poly in enumerate(polys):
        pts = [(E(x + px * k), E(y + (py + 960) * k)) for px, py in poly]
        if j:
            fb.move_to(*pts[0])
        fb.add_line_segments(pts[1:] if j == 0 else pts[1:], close=True)
    s = fb.convert_to_shape()
    _plain(s)
    s.fill.solid()
    s.fill.fore_color.rgb = rgb(color)
    s.line.fill.background()
    s.name = f"아이콘 {name}"
    return s


def _arrow(connector) -> None:
    ln = connector.line._get_or_add_ln()
    tail = etree.SubElement(ln, qn("a:tailEnd"))
    tail.set("type", "triangle")
    tail.set("w", "med")
    tail.set("len", "med")


def connector(slide, kind, p1, p2, color, dashed, name, a=None, a_side=None, b=None, b_side=None,
              arrow=True, width=None):
    c = slide.shapes.add_connector(kind, E(p1[0]), E(p1[1]), E(p2[0]), E(p2[1]))
    _plain(c)
    # 변 가운데(연결점)에 닿는 끝만 붙인다. 비켜난 끝까지 붙이면 뷰어가 선을 연결점으로 다시 그린다.
    def site(shape, side):
        oval = shape.shape_type is not None and getattr(shape, "auto_shape_type", None) == MSO_SHAPE.OVAL
        return (CXN_OVAL if oval else CXN)[side]
    # 위치는 아래에서 직접 정하므로 연결 정보(XML)만 붙인다(python-pptx 이동 계산은 사각형 0~3번만 안다)
    if a is not None:
        c._connect_begin_to(a, site(a, a_side))
    if b is not None:
        c._connect_end_to(b, site(b, b_side))
    c.begin_x, c.begin_y, c.end_x, c.end_y = E(p1[0]), E(p1[1]), E(p2[0]), E(p2[1])
    c.line.color.rgb = rgb(color)
    c.line.width = Pt((width or T["line"]) * PT)
    if dashed:
        c.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    if arrow:
        _arrow(c)
    c.name = name
    return c


def _line(slide, p1, p2, color, width, name, dashed=False):
    return connector(slide, MSO_CONNECTOR.STRAIGHT, p1, p2, color, dashed, name, arrow=False, width=width)


def _marks(slide, marks) -> None:
    for m in marks:
        if m[0] == "line":
            _line(slide, (m[1], m[2]), (m[3], m[4]), C["primary"], 2, "관계 기호")
        else:
            o = slide.shapes.add_shape(MSO_SHAPE.OVAL, E(m[1] - m[3]), E(m[2] - m[3]), E(2 * m[3]), E(2 * m[3]))
            _plain(o)
            o.fill.solid()
            o.fill.fore_color.rgb = rgb(C["white"])
            o.line.color.rgb = rgb(C["primary"])
            o.line.width = Pt(2 * PT)
            o.name = "관계 기호 0"


def _col_box(slide, x, y, w, texts, size, bold, color, align, name):
    """컬럼 한 열(줄 높이 E_ROW 고정). texts 의 빈 글은 빈 줄."""
    tb = slide.shapes.add_textbox(E(x), E(y), E(w), E(E_ROW * len(texts)))
    tf = tb.text_frame
    _margins(tf)
    tf.vertical_anchor = MSO_ANCHOR.TOP
    lines = [(t, size, bold, c, E_ROW) for t, c in texts] if isinstance(texts[0], tuple) else \
        [(t, size, bold, color, E_ROW) for t in texts]
    _paras(tf, lines, align)
    tb.name = name
    return tb


def entity(slide, e: Entity) -> None:
    g = slide.shapes.add_group_shape()
    line = C["block_line"] if e.ref else C["primary"]
    box = rrect(g, e.x, e.y, e.w, e.h, C["white"], line, 1.8, 10, "테두리")
    if e.ref:
        box.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    hd = g.shapes.add_shape(MSO_SHAPE.ROUND_2_SAME_RECTANGLE, E(e.x + 1), E(e.y + 1), E(e.w - 2), E(E_HEAD - 1))
    _plain(hd)
    hd.adjustments[0] = 9 / (E_HEAD - 1)
    hd.adjustments[1] = 0
    hd.fill.solid()
    hd.fill.fore_color.rgb = rgb(C["zone_fill"] if e.ref else C["tint_fill"])
    hd.line.fill.background()
    hd.name = "머리"
    _line(g, (e.x, e.y + E_HEAD), (e.x + e.w, e.y + E_HEAD), line, 1.4, "머리 선")
    textbox(g, e.x + 14, e.y, e.w / 2, E_HEAD, [(e.name, T["block_title"], True, C["navy"], None)], name="테이블 이름")
    textbox(g, e.x + e.w / 2, e.y, e.w / 2 - 14, E_HEAD, [(e.label, T["label"], True, C["muted"], None)],
            PP_ALIGN.RIGHT, name="한글 이름")
    _line(g, (e.x + 10, e.sep_y()), (e.x + e.w - 10, e.sep_y()), C["block_line"], 1.4, "PK 구분선")
    cx = e.col_x()
    f = T["label"]
    rows = e.rows()
    for part, cols in (("PK", rows[:len(e.pk)]), ("컬럼", rows[len(e.pk):])):
        if not cols:
            continue
        y0 = cols[0][0]
        keys = [c[0] for _, c in cols]
        names = [(c[1], C["muted"] if not c[0] and not c[2] else C["text"]) for _, c in cols]
        if any(keys):
            _col_box(g, cx["key"], y0, 34, keys, f, True, C["primary"], PP_ALIGN.LEFT, f"{part} 키")
        _col_box(g, cx["name"], y0, e.w - 200, names, f, part == "PK", None, PP_ALIGN.LEFT, f"{part} 이름")
        _col_box(g, cx["type"] - 150, y0, 150, [c[2] for _, c in cols], f, False, C["muted"], PP_ALIGN.RIGHT, f"{part} 타입")
        _col_box(g, cx["nn"] - 30, y0, 30, ["NN" if c[3] else "" for _, c in cols], f, True, C["muted"],
                 PP_ALIGN.RIGHT, f"{part} NN")
    g.name = f"엔터티 {e.name}"


def relation(slide, r: Rel) -> None:
    g = slide.shapes.add_group_shape()
    fb = g.shapes.build_freeform(E(r.pts[0][0]), E(r.pts[0][1]), scale=1.0)
    fb.add_line_segments([(E(x), E(y)) for x, y in r.pts[1:]], close=False)
    s = fb.convert_to_shape()
    _plain(s)
    s.fill.background()
    s.line.color.rgb = rgb(C["primary"])
    s.line.width = Pt(2 * PT)
    if not r.ident:
        s.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    s.name = "관계선"
    _marks(g, ie_marks(r.pts[0], r.pts[1], r.parent_card))
    _marks(g, ie_marks(r.pts[-1], r.pts[-2], r.child_card))
    g.name = f"관계 {r.parent} → {r.child}"


def build(ds, out_path, title: str) -> None:
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(SLIDE_W), Emu(SLIDE_H)
    for d in (ds if isinstance(ds, list) else [ds]):
        _slide(prs, d)
    prs.core_properties.title = title
    prs.core_properties.author = "KDMS"
    prs.save(out_path)


def _slide(prs, d: Diagram) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = rgb("#FFFFFF")

    # 머리글
    bar = rrect(slide, 50, 36, 7, 72, C["primary"], None, 0, 3.5, "제목 막대")
    textbox(slide, 70, 30, 1050, 48, [(d.title, T["title"], True, C["navy"], None)], name="제목")
    textbox(slide, 70, 84, 1200, 24, [(d.subtitle, T["subtitle"], False, C["muted"], None)], name="부제")
    textbox(slide, 1100, 54, 450, 22, [(d.meta, T["meta"], False, C["muted"], None)], PP_ALIGN.RIGHT, name="문서 정보")
    del bar

    # 영역(제목·부제는 도형 안)
    for z in d.zones:
        s = rrect(slide, z.x, z.y, z.w, z.h,
                  C["core_zone_fill"] if z.core else C["zone_fill"],
                  C["core_zone_line"] if z.core else C["zone_line"], 2 if z.core else 1.4, 18, f"영역 {z.title}")
        tf = s.text_frame
        _margins(tf, 22, 14, 22, 0)
        tf.vertical_anchor = MSO_ANCHOR.TOP
        lines = [(z.title, T["zone_title"], True, C["navy"], 26)]
        for sub in (z.sub.split("\n") if z.sub else []):
            lines.append((sub, T["zone_sub"], False, C["muted"], T["zone_sub_line"]))
        _paras(tf, lines, PP_ALIGN.LEFT if z.align == "start" else PP_ALIGN.RIGHT)

    # 막대·자유 선(블록 아래에 깔린다)
    for b in d.bars:
        s = rrect(slide, b.x, b.y, b.w, b.h, b.fill, b.line, b.lw, b.rx, b.name)
        if b.text:
            tf = s.text_frame
            _margins(tf)
            tf.vertical_anchor = MSO_ANCHOR.MIDDLE
            _paras(tf, [(b.text, b.size, True, b.text_color, None)], PP_ALIGN.CENTER)
    for l in d.lines:
        connector(slide, MSO_CONNECTOR.STRAIGHT, (l.x1, l.y1), (l.x2, l.y2), l.color, l.dashed, "선",
                  arrow=l.arrow, width=l.width)

    # ERD 관계선 → 엔터티(기호가 엔터티 밖에 있어 순서는 상관없다)
    for r in d.rels:
        relation(slide, r)
    for e in d.entities.values():
        entity(slide, e)

    # 블록
    shapes = {}
    for b in d.blocks.values():
        st = b.style()
        if b.pill or b.shape == "start":
            dark = b.shape == "start"
            s = rrect(slide, b.x, b.y, b.w, b.h, C["navy"] if dark else C["white"],
                      C["navy"] if dark else C["primary"], 2, b.h / 2, f"블록 {b.title}")
            ix, tx = d.pill_layout(b)
            tf = s.text_frame
            _margins(tf, tx - b.x, 0, 0, 0)
            tf.vertical_anchor = MSO_ANCHOR.MIDDLE
            _paras(tf, [(b.title, T["pill"], True, C["white"] if dark else C["navy"], None)], PP_ALIGN.LEFT)
            if b.icon:
                icon(slide, b.icon, ix, b.y + (b.h - T["pill_icon"]) / 2, T["pill_icon"],
                     C["white"] if dark else C["primary"])
            shapes[b.id] = s
            continue
        if b.shape == "point":
            s = slide.shapes.add_shape(MSO_SHAPE.OVAL, E(b.x), E(b.y), E(b.w), E(b.h))
            _plain(s)
            s.fill.solid()
            s.fill.fore_color.rgb = rgb(st["fill"])
            s.line.color.rgb = rgb(st["line"])
            s.line.width = Pt(st["lw"] * PT)
            tf = s.text_frame
            _margins(tf)
            tf.vertical_anchor = MSO_ANCHOR.MIDDLE
            _paras(tf, [(b.title, T["block_title"], True, st["title"], None)], PP_ALIGN.CENTER)
            s.name = f"연결점 {b.title}"
            shapes[b.id] = s
            continue
        if b.shape == "diamond":
            s = slide.shapes.add_shape(MSO_SHAPE.DIAMOND, E(b.x), E(b.y), E(b.w), E(b.h))
            _plain(s)
            s.fill.solid()
            s.fill.fore_color.rgb = rgb(st["fill"])
            s.line.color.rgb = rgb(st["line"])
            s.line.width = Pt(st["lw"] * PT)
            s.name = f"판단 {b.title}"
        else:
            s = rrect(slide, b.x, b.y, b.w, b.h, st["fill"], st["line"], st["lw"], 12, f"블록 {b.title}")
        lay = b.layout()
        tf = s.text_frame
        _margins(tf, 6, lay["text_top"] - b.y - 4, 6, 0)  # 뷰어 줄간격 여유
        tf.vertical_anchor = MSO_ANCHOR.TOP
        lines = [(b.title, T["block_title"], True, st["title"], T["title_line"])]
        lines += [(x, T["block_sub"], False, st["sub"], T["sub_line"]) for x in b.sub]
        _paras(tf, lines, PP_ALIGN.CENTER)
        if b.icon:
            icon(slide, b.icon, lay["cx"] - T["icon"] / 2, lay["icon_y"], T["icon"], st["icon"])
        if b.num:
            r = T["badge_r"]
            o = slide.shapes.add_shape(MSO_SHAPE.OVAL, E(b.x + 21 - r), E(b.y + 21 - r), E(2 * r), E(2 * r))
            _plain(o)
            o.fill.solid()
            o.fill.fore_color.rgb = rgb(st["badge"])
            o.line.fill.background()
            tf = o.text_frame
            _margins(tf)
            tf.vertical_anchor = MSO_ANCHOR.MIDDLE
            _paras(tf, [(b.num, T["badge"], True, st["badge_text"], None)], PP_ALIGN.CENTER)
            o.name = f"번호 {b.num}"
        shapes[b.id] = s

    # 연결선과 라벨
    for l in d.links:
        pts = d.route(l)
        kind = MSO_CONNECTOR.STRAIGHT if len(pts) == 2 else MSO_CONNECTOR.ELBOW
        color = C["sky"] if l.dashed else C["primary"]
        connector(slide, kind, pts[0], pts[-1], color, l.dashed, f"연결 {l.a}→{l.b}",
                  shapes[l.a] if not l.a_off else None, l.a_side,
                  shapes[l.b] if not l.b_off else None, l.b_side)
        if l.label:
            x, y, w, h, _, _ = d.label_box(l)
            textbox(slide, x, y, w, h, [(l.label, T["label"], True, color, None)], PP_ALIGN.CENTER,
                    name=f"라벨 {l.label}", fill=C["white"])

    # 자유 글(기준선 → 글상자 위치)
    for t in d.texts:
        w = text_width(t.text, t.size) + 24
        x = {"start": t.x, "middle": t.x - w / 2, "end": t.x - w}[t.anchor]
        al = {"start": PP_ALIGN.LEFT, "middle": PP_ALIGN.CENTER, "end": PP_ALIGN.RIGHT}[t.anchor]
        textbox(slide, x, t.y - t.size - 2, w, t.size + 8, [(t.text, t.size, t.bold, t.color, None)], al, name="글")

    # 범례(왼쪽 아래)
    if d.legend:
        (x, y, w, h), items = d.legend_layout()
        my = y + h / 2
        rrect(slide, x, y, w, h, C["white"], C["zone_line"], 1.4, 10, "범례")
        textbox(slide, x + 18, y, 50, h, [("범례", T["legend"], True, C["navy"], None)], name="범례 제목")
        for kind, label, mx, tx in items:
            if kind in ("sync", "async"):
                connector(slide, MSO_CONNECTOR.STRAIGHT, (mx, my), (mx + 46, my),
                          C["primary"] if kind == "sync" else C["sky"], kind == "async", f"범례 {label}")
            elif kind in LINE_KINDS:
                _line(slide, (mx, my), (mx + 46, my), C["primary"], 2, f"범례 {label}", dashed=kind == "nonident")
                if kind.startswith("ie_"):
                    _marks(slide, ie_marks((mx + 46, my), (mx, my), kind[3:]))
            else:
                rrect(slide, mx, my - 9, 24, 18, kind, C["block_line"], 1, 4, f"범례 {label}")
            textbox(slide, tx, y, text_width(label, T["legend"]) + 20, h,
                    [(label, T["legend"], True, C["text"], None)], name=f"범례 글 {label}")

    for nx, ny, n in d.note_positions():
        wdt = text_width(n, T["note"]) + 40
        textbox(slide, nx - wdt, ny - T["note"] - 2, wdt, T["note"] + 8,
                [(n, T["note"], False, C["muted"], None)], PP_ALIGN.RIGHT, name="설명")

